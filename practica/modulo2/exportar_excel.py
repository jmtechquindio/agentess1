#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Exporta los CSV validados a un libro de Excel (datos_modulo2.xlsx)."""

import csv
import collections
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent


def leer(nombre):
    with (ROOT / nombre).open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


conv = leer("conversaciones.csv")
lineas = leer("lineas_producto.csv")
pend = leer("pendientes_excepciones.csv")

wb = Workbook()

HEADER_FILL = PatternFill("solid", fgColor="14416B")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ESTADO_FILL = {
    "confirmado": PatternFill("solid", fgColor="DDEFE6"),
    "pendiente": PatternFill("solid", fgColor="FDF3DC"),
    "cancelado": PatternFill("solid", fgColor="FBE3E3"),
    "sin_pedido": PatternFill("solid", fgColor="E9EDF2"),
}


def hoja(ws, rows, fills=None):
    if not rows:
        return
    cols = list(rows[0].keys())
    ws.append(cols)
    for c in cols:
        cell = ws.cell(row=1, column=cols.index(c) + 1)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
        if fills:
            est = r.get("estado")
            color = fills.get(est)
            if color:
                for i in range(1, len(cols) + 1):
                    ws.cell(row=ws.max_row, column=i).fill = color
    for i, c in enumerate(cols, 1):
        width = max(len(c) + 2, max((len(str(r.get(c, ""))) for r in rows), default=4) + 1)
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    ws.freeze_panes = "A2"


# Hoja 1: conversaciones
ws1 = wb.active
ws1.title = "Conversaciones"
hoja(ws1, conv, fills=ESTADO_FILL)

# Hoja 2: lineas de producto
ws2 = wb.create_sheet("Lineas_Producto")
hoja(ws2, lineas)

# Hoja 3: pendientes / excepciones
ws3 = wb.create_sheet("Pendientes_Excepciones")
hoja(ws3, pend)

# Hoja 4: resumen
ws4 = wb.create_sheet("Resumen")
estados = collections.Counter(x["estado"] for x in conv)
resumen = [
    {"concepto": "Total conversaciones", "valor": len(conv)},
    {"concepto": "Pedidos confirmados", "valor": estados["confirmado"]},
    {"concepto": "Sin pedido (consultas)", "valor": estados["sin_pedido"]},
    {"concepto": "Pendientes", "valor": estados["pendiente"]},
    {"concepto": "Cancelados", "valor": estados["cancelado"]},
    {"concepto": "Líneas de producto", "valor": len(lineas)},
    {"concepto": "Pendientes / excepciones", "valor": len(pend)},
    {"concepto": "Total ventas confirmados ($)", "valor": sum(int(x["total"]) for x in conv if x["estado"] == "confirmado" and x["total"])},
]
hoja(ws4, resumen)

out = ROOT / "datos_modulo2.xlsx"
wb.save(out)
print("generado:", out, f"({out.stat().st_size} bytes)")